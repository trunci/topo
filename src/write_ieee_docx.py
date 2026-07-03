"""Render paper/PAPER.md into the IEEE conference (letter) Word template.

Same provenance discipline as the rest of the pipeline: PAPER.md is generated
from the stats JSONs by write_paper.py; this script only reformats it, so no
number is ever hand-typed. The IEEE template (paper/ieee_template.docx) is the
unmodified conference-template-letter.docx distributed by IEEE.

Layout notes:
* title / author / provenance note span the full page width (single-column
  section), everything after flows in the template's two-column body section;
* the template's Heading1-3 auto-numbering (I., A., ...) is suppressed because
  the paper's own section numbers ("5.3") are load-bearing — in-text
  cross-references (§5.3 etc.) must keep resolving;
* citations stay author-year (not IEEE numeric); references are emitted as
  [n]-prefixed paragraphs in citation order of the References section.

Regenerate with: uv run python -m src.write_ieee_docx
"""
from __future__ import annotations

import copy
import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt, Inches

from src.ieee_template_convert import strict_to_transitional

PAPER_MD = Path("paper/PAPER.md")
TEMPLATE = Path("paper/ieee_template.docx")
OUT = Path("paper/PAPER_ieee.docx")
COLUMN_WIDTH_IN = 3.5  # two-column IEEE letter column


# ---------------------------------------------------------------- md parsing

def parse_blocks(md: str):
    """Yield (kind, payload) blocks: title/author/note/h1/h2/para/table/img/
    caption/listitem/refitem."""
    lines = md.split("\n")
    i, blocks = 0, []
    para: list[str] = []

    def flush():
        nonlocal para
        if para:
            blocks.append(("para", " ".join(para)))
            para = []

    while i < len(lines):
        ln = lines[i]
        if ln.startswith("# ") and not blocks:
            blocks.append(("title", ln[2:].strip()))
        elif re.fullmatch(r"\*\*[^*]+\*\*", ln.strip()) and len(blocks) == 1:
            blocks.append(("author", ln.strip().strip("*")))
        elif ln.startswith("*All numbers"):
            note = [ln]
            while not note[-1].rstrip().endswith("*") or len(note) == 1 \
                    and not note[0].rstrip().endswith(".*"):
                i += 1
                note.append(lines[i])
                if note[-1].rstrip().endswith("*"):
                    break
            blocks.append(("note", " ".join(x.strip() for x in note).strip("*")))
        elif ln.startswith("## "):
            flush()
            blocks.append(("h1", ln[3:].strip()))
        elif ln.startswith("### "):
            flush()
            blocks.append(("h2", ln[4:].strip()))
        elif ln.startswith("|"):
            flush()
            tbl = []
            while i < len(lines) and lines[i].startswith("|"):
                cells = [c.strip() for c in lines[i].strip("|").split("|")]
                if not all(re.fullmatch(r":?-+:?", c) for c in cells):
                    tbl.append(cells)
                i += 1
            i -= 1
            blocks.append(("table", tbl))
        elif ln.startswith("!["):
            m = re.match(r"!\[[^\]]*\]\(([^)]+)\)", ln)
            blocks.append(("img", m.group(1)))
        elif re.match(r"\*Figure \d+:", ln):
            cap = [ln]
            while not cap[-1].rstrip().endswith("*"):
                i += 1
                cap.append(lines[i])
            blocks.append(("caption", " ".join(x.strip() for x in cap).strip("*")))
        elif re.match(r"\d+\. \*\*", ln):
            item = [ln]
            while i + 1 < len(lines) and lines[i + 1].startswith("   "):
                i += 1
                item.append(lines[i].strip())
            blocks.append(("listitem", " ".join(x.strip() for x in item)))
        elif ln.startswith("- "):
            item = [ln[2:]]
            while i + 1 < len(lines) and lines[i + 1].startswith("  ") \
                    and not lines[i + 1].startswith("- "):
                i += 1
                item.append(lines[i].strip())
            blocks.append(("refitem", " ".join(item)))
        elif ln.strip() == "":
            flush()
        else:
            para.append(ln.strip())
        i += 1
    flush()
    return blocks


INLINE = re.compile(r"(\*\*.+?\*\*|\*[^*]+?\*|`[^`]+?`)")


def add_runs(paragraph, text, size=None):
    for tok in INLINE.split(text):
        if not tok:
            continue
        if tok.startswith("**") and tok.endswith("**"):
            r = paragraph.add_run(tok[2:-2])
            r.bold = True
        elif tok.startswith("*") and tok.endswith("*") and len(tok) > 2:
            r = paragraph.add_run(tok[1:-1])
            r.italic = True
        elif tok.startswith("`") and tok.endswith("`"):
            r = paragraph.add_run(tok[1:-1])
            r.font.name = "Courier New"
        else:
            r = paragraph.add_run(tok)
        if size is not None:
            r.font.size = Pt(size)


# ---------------------------------------------------------------- docx build

def suppress_numbering(paragraph):
    """Override the template's Heading auto-numbering (numId=0)."""
    pPr = paragraph._p.get_or_add_pPr()
    numPr = pPr.makeelement(qn("w:numPr"), {})
    ilvl = pPr.makeelement(qn("w:ilvl"), {qn("w:val"): "0"})
    numId = pPr.makeelement(qn("w:numId"), {qn("w:val"): "0"})
    numPr.append(ilvl)
    numPr.append(numId)
    # OOXML requires numPr immediately after pStyle within pPr
    pStyle = pPr.find(qn("w:pStyle"))
    if pStyle is not None:
        pStyle.addnext(numPr)
    else:
        pPr.insert(0, numPr)


def set_columns(sectPr, num, space_pt):
    cols = sectPr.find(qn("w:cols"))
    if cols is None:
        cols = sectPr.makeelement(qn("w:cols"), {})
        sectPr.append(cols)
    if num > 1:
        cols.set(qn("w:num"), str(num))
    elif cols.get(qn("w:num")):
        del cols.attrib[qn("w:num")]
    cols.set(qn("w:space"), str(space_pt * 20))  # twips


def add_table_borders(table):
    tblPr = table._tbl.tblPr
    borders = tblPr.makeelement(qn("w:tblBorders"), {})
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = borders.makeelement(qn(f"w:{edge}"), {
            qn("w:val"): "single", qn("w:sz"): "4",
            qn("w:space"): "0", qn("w:color"): "999999"})
        borders.append(el)
    tblPr.append(borders)


def main():
    blocks = parse_blocks(PAPER_MD.read_text())
    doc = Document(strict_to_transitional(str(TEMPLATE)))

    # wipe the template's sample content (paragraphs and tables), keep styles
    body = doc.element.body
    for el in list(body):
        if el.tag in (qn("w:p"), qn("w:tbl")):
            body.remove(el)

    # final (body) section: two columns, IEEE gutter
    final_sectPr = body.find(qn("w:sectPr"))
    set_columns(final_sectPr, 2, 18)

    title_sectPr = copy.deepcopy(final_sectPr)  # single-column title block
    set_columns(title_sectPr, 1, 36)

    ref_counter = 0
    current_h1 = None
    for kind, payload in blocks:
        if kind == "title":
            p = doc.add_paragraph(style="paper title")
            add_runs(p, payload)
        elif kind == "author":
            p = doc.add_paragraph(style="Author")
            add_runs(p, payload)
        elif kind == "note":
            p = doc.add_paragraph(style="Body Text")
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(payload)
            r.italic = True
            r.font.size = Pt(8)
            # close the single-column title-block section
            brk = doc.add_paragraph()
            brk._p.get_or_add_pPr().append(title_sectPr)
        elif kind == "h1" and payload == "Abstract":
            current_h1 = payload  # folded into the Abstract-styled paragraph
        elif kind == "h1":
            current_h1 = payload
            p = doc.add_paragraph(style="Heading 1")
            suppress_numbering(p)
            add_runs(p, payload)
        elif kind == "h2":
            p = doc.add_paragraph(style="Heading 2")
            suppress_numbering(p)
            add_runs(p, payload)
        elif kind == "para":
            if current_h1 == "Abstract":
                p = doc.add_paragraph(style="Abstract")
                r = p.add_run("Abstract—")
                r.bold = True
                add_runs(p, payload)
            else:
                p = doc.add_paragraph(style="Body Text")
                add_runs(p, payload)
        elif kind == "listitem":
            p = doc.add_paragraph(style="Body Text")
            add_runs(p, payload)
        elif kind == "refitem":
            ref_counter += 1
            p = doc.add_paragraph(style="Body Text")
            p.paragraph_format.first_line_indent = Inches(0)
            add_runs(p, f"[{ref_counter}] {payload}", size=8)
        elif kind == "img":
            fig_counter += 1
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.add_run().add_picture(str(PAPER_MD.parent / payload),
                                    width=Inches(COLUMN_WIDTH_IN))
        elif kind == "caption":
            p = doc.add_paragraph(style="figure caption")
            add_runs(p, payload)
        elif kind == "table":
            header, *rows = payload
            t = doc.add_table(rows=len(payload), cols=len(header))
            t.alignment = WD_TABLE_ALIGNMENT.CENTER
            add_table_borders(t)
            for ci, cell in enumerate(header):
                para = t.rows[0].cells[ci].paragraphs[0]
                r = para.add_run(cell)
                r.bold = True
                r.font.size = Pt(8)
            for ri, row in enumerate(rows, start=1):
                for ci, cell in enumerate(row):
                    para = t.rows[ri].cells[ci].paragraphs[0]
                    add_runs(para, cell, size=8)
            doc.add_paragraph(style="Body Text")  # spacing after table

    doc.save(str(OUT))
    print(f"wrote {OUT}")


if __name__ == "__main__":
    main()
